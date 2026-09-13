# -*- coding: utf-8 -*-
"""Excel 导出：竞品数据明细 + 横向对比 + 词频统计

表头设计以“便于二次编辑”为原则：
- 明细表每行一条商品，列字段固定、表头加色加粗、首行冻结、开启筛选
- 对比表把同一搜索词下的商品横向排开，方便直接看卖点/成分差异
- 词频表汇总高频功效/成分/关键词
"""
import io
import re
from collections import Counter
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    ("序号", 6),
    ("平台", 10),
    ("商品名称", 34),
    ("品牌", 14),
    ("关键词", 28),
    ("功效", 26),
    ("核心卖点", 46),
    ("成分", 34),
    ("成分差异", 26),
    ("价格(元)", 11),
    ("销量", 11),
    ("销量原文", 13),
    ("店铺/作者", 20),
    ("商品链接", 42),
    ("搜索词/分组", 14),
    ("数据来源", 10),
    ("抓取时间", 18),
    ("备注", 18),
]

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10)
LINK_FONT = Font(name="微软雅黑", size=10, color="0563C1", underline="single")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
CENTER_WRAP = Alignment(horizontal="center", vertical="top", wrap_text=True)


def _split(text):
    if not text:
        return []
    return [t.strip() for t in re.split(r"[、，,；;｜|/]", str(text)) if t.strip()]


def _counter(items):
    c = Counter()
    for it in items:
        for w in _split(it):
            c[w] += 1
    return c.most_common(10)


def _fmt_num(v, fmt="#,##0.##"):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def build_workbook(items, meta=None):
    meta = meta or {}
    wb = Workbook()

    # ---------------------------------------------------------- Sheet1 明细
    ws = wb.active
    ws.title = "竞品数据明细"
    ws.append([h[0] for h in HEADERS])
    for col, (_, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for idx, it in enumerate(items, start=1):
        row = [
            idx,
            it.get("platform") or "",
            it.get("name") or "",
            it.get("brand") or "",
            it.get("keywords") or "",
            it.get("effects") or "",
            it.get("selling_points") or "",
            it.get("ingredients") or "",
            it.get("ingredient_diff") or "",
            _fmt_num(it.get("price")),
            _fmt_num(it.get("sales")),
            it.get("sales_text") or "",
            it.get("shop") or "",
            it.get("url") or "",
            it.get("keyword_group") or "",
            it.get("source") or "",
            it.get("created_at") or "",
            it.get("note") or "",
        ]
        ws.append(row)
        r = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = BORDER
            if col in (10, 11):
                cell.number_format = "#,##0.##"
                cell.alignment = CENTER
            elif col in (14,):
                cell.alignment = LEFT_WRAP
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = LINK_FONT
            elif col in (6, 7, 8, 9):
                cell.alignment = LEFT_WRAP
            else:
                cell.alignment = CENTER_WRAP if col in (1, 2, 12, 13, 15, 16, 17, 18) else LEFT_WRAP
            cell.font = BODY_FONT
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # ---------------------------------------------------------- Sheet2 对比
    ws2 = wb.create_sheet("横向对比")
    groups = {}
    for it in items:
        groups.setdefault(it.get("keyword_group") or "全部商品", []).append(it)
    compare_rows = [
        ("平台", "platform", CENTER_WRAP),
        ("品牌", "brand", LEFT_WRAP),
        ("价格(元)", "price", CENTER_WRAP),
        ("销量", "sales", CENTER_WRAP),
        ("功效", "effects", LEFT_WRAP),
        ("核心卖点", "selling_points", LEFT_WRAP),
        ("成分", "ingredients", LEFT_WRAP),
        ("成分差异", "ingredient_diff", LEFT_WRAP),
        ("商品链接", "url", LEFT_WRAP),
        ("备注", "note", LEFT_WRAP),
    ]
    row = 1
    for group, its in groups.items():
        start = row
        ws2.cell(row=row, column=1, value=f"对比组：{group}（{len(its)} 个商品）")
        ws2.cell(row=row, column=1).font = Font(name="微软雅黑", size=11, bold=True)
        ws2.merge_cells(start_row=row, start_column=1, end_row=row, end_column=1 + len(its))
        row += 1
        ws2.cell(row=row, column=1, value="对比维度")
        for j, it in enumerate(its, start=2):
            ws2.cell(row=row, column=j, value=(it.get("name") or "商品")[:40])
        for j in range(1, len(its) + 2):
            cell = ws2.cell(row=row, column=j)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER_WRAP
            cell.border = BORDER
        row += 1
        for label, key, align in compare_rows:
            ws2.cell(row=row, column=1, value=label)
            for j, it in enumerate(its, start=2):
                val = it.get(key) or ""
                cell = ws2.cell(row=row, column=j, value=val)
                if key in ("price", "sales") and val != "":
                    cell.number_format = "#,##0.##"
                if key == "url" and val:
                    cell.hyperlink = str(val)
                    cell.font = LINK_FONT
                cell.alignment = align
                cell.border = BORDER
                cell.font = BODY_FONT
            ws2.cell(row=row, column=1).border = BORDER
            ws2.cell(row=row, column=1).font = Font(name="微软雅黑", size=10, bold=True)
            row += 1
        row += 1  # 空行分隔
    ws2.column_dimensions["A"].width = 16
    for j in range(2, max(2, len(items) + 2) + 1):
        ws2.column_dimensions[get_column_letter(j)].width = 38
    ws2.freeze_panes = "B2"

    # ---------------------------------------------------------- Sheet3 词频
    ws3 = wb.create_sheet("词频统计")
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 40
    ws3.column_dimensions["C"].width = 10
    ws3.column_dimensions["D"].width = 16
    ws3.column_dimensions["E"].width = 40
    ws3.column_dimensions["F"].width = 10

    def freq_table(r0, title, pairs):
        ws3.cell(row=r0, column=1, value=title).font = Font(name="微软雅黑", size=11, bold=True)
        ws3.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=3)
        r0 += 1
        for c, t in ((1, "类别"), (2, "词"), (3, "出现次数")):
            cell = ws3.cell(row=r0, column=c, value=t)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = CENTER
            cell.border = BORDER
        r0 += 1
        for label, w, n in pairs:
            ws3.cell(row=r0, column=1, value=label)
            ws3.cell(row=r0, column=2, value=w)
            ws3.cell(row=r0, column=3, value=n)
            for c in (1, 2, 3):
                cell = ws3.cell(row=r0, column=c)
                cell.border = BORDER
                cell.font = BODY_FONT
                cell.alignment = LEFT_WRAP if c == 2 else CENTER
            r0 += 1
        return r0 + 1

    r = 1
    r = freq_table(r, "高频功效 TOP10", [("功效", w, n) for w, n in _counter([it.get("effects") for it in items])])
    r = freq_table(r, "高频成分 TOP10", [("成分", w, n) for w, n in _counter([it.get("ingredients") for it in items])])
    r = freq_table(r, "高频关键词 TOP10", [("关键词", w, n) for w, n in _counter([it.get("keywords") for it in items])])

    # 汇总信息
    r += 1
    ws3.cell(row=r, column=1, value="汇总").font = Font(name="微软雅黑", size=11, bold=True)
    r += 1
    info = [
        ("商品总数", len(items)),
        ("涉及平台", "、".join(sorted({it.get("platform") or "" for it in items}))),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("搜索词", meta.get("keyword") or "全部"),
    ]
    for k, v in info:
        ws3.cell(row=r, column=1, value=k).font = Font(name="微软雅黑", size=10, bold=True)
        ws3.cell(row=r, column=2, value=v).font = BODY_FONT
        r += 1

    return wb


def export_xlsx(items, meta=None):
    """返回 xlsx 字节流"""
    wb = build_workbook(items, meta)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_filename(keyword="", platform=""):
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    kw = re.sub(r"[\\/:*?\"<>|]", "_", keyword or "")[:20]
    pf = re.sub(r"[\\/:*?\"<>|]", "_", platform or "")
    parts = [p for p in ("竞品分析", kw, pf) if p]
    return "_".join(parts) + f"_{ts}.xlsx"
